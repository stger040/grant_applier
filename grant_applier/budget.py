"""Budget workbook and assumptions generation."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from .discovery import Opportunity
from .profiles import ensure_profile


def generate_budget(opportunity: Opportunity, overwrite: bool = False) -> list[Path]:
    profile = ensure_profile(opportunity, overwrite=False)
    budgets_dir = opportunity.path / "Budgets"
    budgets_dir.mkdir(parents=True, exist_ok=True)

    created: list[Path] = []
    main_budget_path = budgets_dir / "applicant_budget.xlsx"
    if overwrite or not main_budget_path.exists():
        write_budget_workbook(main_budget_path, profile, include_subaward=False)
        created.append(main_budget_path)

    if opportunity.applicant_folder == "AAIP":
        subaward_path = budgets_dir / "subaward_budget.xlsx"
        if overwrite or not subaward_path.exists():
            write_budget_workbook(subaward_path, profile, include_subaward=True)
            created.append(subaward_path)

    assumptions_path = budgets_dir / "budget_assumptions.md"
    if overwrite or not assumptions_path.exists():
        assumptions_path.write_text(render_budget_assumptions(profile).rstrip() + "\n", encoding="utf-8")
        created.append(assumptions_path)
    return created


def write_budget_workbook(path: Path, profile: dict[str, Any], include_subaward: bool) -> None:
    try:
        from openpyxl import Workbook
    except ImportError as error:
        raise RuntimeError(
            "openpyxl is required for .xlsx budget generation. Install with: python3 -m pip install openpyxl"
        ) from error

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    assumptions = workbook.create_sheet("Assumptions")
    categories = workbook.create_sheet("Line_Items")

    summary["A1"] = "Opportunity Number"
    summary["B1"] = profile.get("opportunity_number", "[CONFIRM]")
    summary["A2"] = "Applicant"
    summary["B2"] = profile.get("likely_applicant", "[CONFIRM]")
    summary["A3"] = "Funder"
    summary["B3"] = profile.get("funder", "[CONFIRM]")
    summary["A4"] = "Status"
    summary["B4"] = "Draft budget template; not submission-ready."
    summary["A6"] = "Year 1 Total"
    summary["A7"] = "Year 2 Total"
    summary["A8"] = "Year 3 Total"
    summary["A9"] = "Total Project Budget"
    summary["B6"] = "=SUM(Line_Items!F2:F20)"
    summary["B7"] = "=SUM(Line_Items!G2:G20)"
    summary["B8"] = "=SUM(Line_Items!H2:H20)"
    summary["B9"] = "=SUM(B6:B8)"

    assumptions["A1"] = "Assumption"
    assumptions["B1"] = "Value"
    assumptions["A2"] = "Fringe rate"
    assumptions["B2"] = "[CONFIRM]"
    assumptions["A3"] = "Indirect cost rate"
    assumptions["B3"] = "[CONFIRM]"
    assumptions["A4"] = "Cost share required"
    assumptions["B4"] = profile.get("cost_share_required", "[CONFIRM]")
    assumptions["A5"] = "Subaward included"
    assumptions["B5"] = "Yes" if include_subaward else "No"

    headers = [
        "Category",
        "Description",
        "Quantity",
        "Unit Cost",
        "Calculation",
        "Year 1",
        "Year 2",
        "Year 3",
    ]
    for index, header in enumerate(headers, start=1):
        categories.cell(row=1, column=index).value = header

    starter_rows = [
        ("Personnel", "Project Director", 1, "[CONFIRM]", "", "[CONFIRM]", "[CONFIRM]", "[CONFIRM]"),
        ("Personnel", "Program Manager", 1, "[CONFIRM]", "", "[CONFIRM]", "[CONFIRM]", "[CONFIRM]"),
        ("Fringe", "Fringe Benefits", 1, "[CONFIRM]", "", "[CONFIRM]", "[CONFIRM]", "[CONFIRM]"),
        ("Travel", "Project travel and convenings", 1, "[CONFIRM]", "", "[CONFIRM]", "[CONFIRM]", "[CONFIRM]"),
        ("Supplies", "Operational supplies", 1, "[CONFIRM]", "", "[CONFIRM]", "[CONFIRM]", "[CONFIRM]"),
        ("Contractual", "Technical and service contracts", 1, "[CONFIRM]", "", "[CONFIRM]", "[CONFIRM]", "[CONFIRM]"),
        ("Other", "Software, cloud, data systems", 1, "[CONFIRM]", "", "[CONFIRM]", "[CONFIRM]", "[CONFIRM]"),
        ("Indirect", "Indirect costs", 1, "[CONFIRM]", "", "[CONFIRM]", "[CONFIRM]", "[CONFIRM]"),
    ]
    if include_subaward:
        starter_rows.append(
            ("Subaward", "Technical implementation subrecipient", 1, "[CONFIRM]", "", "[CONFIRM]", "[CONFIRM]", "[CONFIRM]")
        )
    for row_index, row in enumerate(starter_rows, start=2):
        for column_index, value in enumerate(row, start=1):
            categories.cell(row=row_index, column=column_index).value = value

    workbook.save(path)


def render_budget_assumptions(profile: dict[str, Any]) -> str:
    return f"""# Budget Assumptions

Opportunity: `{profile.get("opportunity_number", "[CONFIRM]")}`
Applicant: `{profile.get("likely_applicant", "[CONFIRM]")}`

## Provisional Values Requiring Confirmation

- [CONFIRM] Salary rates and effort allocations by role.
- [CONFIRM] Fringe rate and basis.
- [CONFIRM] Indirect cost rate agreement and applicability.
- [CONFIRM] Allowability of cloud/software/consultant categories.
- [CONFIRM] Match or cost share requirement status.
- [CONFIRM] Subaward structure and amount, if applicable.

## Controls

- Do not submit budget forms until all `[CONFIRM]` fields are replaced with validated values.
- Ensure budget totals remain within any published ceiling.
- Ensure budget narrative language aligns with official funder requirements.
"""
